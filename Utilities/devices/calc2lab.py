from docx import Document
from math import *
doc = Document()
table = doc.add_table(rows=4, cols=6)
table.cell(0, 0).text = "Гармоника"
table.cell(1, 0).text = "Эксперимент"
table.cell(2, 0).text = "Расчет"
table.cell(3, 0).text = "Ошибка"

f = "y = "
exps = [1.221, 0.399, 0.235, 0.165, 0.126]
for i in range(1, 6):
    table.cell(0, i).text = f"{2*i-1}"
    table.cell(1, i).text = f"{exps[i-1]}"
    Val = 4/pi/(2*i-1)
    table.cell(2, i).text = f"{round(Val, 3)}"
    table.cell(3, i).text = f"{round(abs(exps[i-1] - Val), 3)}"
    f = f + f"{round(Val, 3)} * sin{2*i-1}x + "
doc.save("4lab.docx")
print(f)