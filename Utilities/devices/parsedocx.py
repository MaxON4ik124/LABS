from docx import Document
from docx.enum.shape import WD_INLINE_SHAPE_TYPE
from docx.shared import Pt, Inches, Cm
from docx.enum.dml import MSO_COLOR_TYPE
doc1 = Document()
doc = Document('bilety.docx')
i = 0
for paragraph in doc.paragraphs:
    doc1.add_paragraph()
    for run in paragraph.runs:
        if run.font.color.type != MSO_COLOR_TYPE.THEME:
            run_copy = doc1.paragraphs[i].add_run(run.text)
            run_copy.font.size = Pt(6)
    doc1.paragraphs[i].paragraph_format.space_before = Pt(0)      
    doc1.paragraphs[i].paragraph_format.space_after = Pt(0)
    i += 1

for paragraph in doc1.paragraphs:
    if len(paragraph.runs) > 0:
        if paragraph.runs[0].text[1:3] == ". " or paragraph.runs[0].text[2:4] == ". ":
            paragraph.paragraph_format.space_before = Pt(6)
    
img = 0















doc1.save("A.docx")
    