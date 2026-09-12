from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_COLOR_INDEX
import re
name = input("Введите полное название файла: ")
Yellow = WD_COLOR_INDEX.YELLOW
Green = WD_COLOR_INDEX.BRIGHT_GREEN
Cyan = WD_COLOR_INDEX.TURQUOISE
doc = Document(name)
Cfind = False
if doc.tables[0]:
    table = doc.tables[0]
    for row in table.rows[1:]:
        cell = row.cells[1]
        Ytext, Gtext, Ctext = [], [], []
        for para in cell.paragraphs:
            runs = enumerate(para.runs)
            for i, r in runs:
                if r.font.bold:
                    r.font.bold = False
                    r.font.highlight_color = Yellow
                    if r.text.strip():
                        Ytext.append(r.text)
                if r.font.italic and r.text[0] == '$' and r.text[-1] == '$':
                    r.font.italic = False
                    r.text = r.text[1:-1]
                    r.font.highlight_color = Cyan
                    if r.text.strip():
                        Ctext.append(r.text)
                if r.font.italic:
                    r.font.italic = False
                    r.font.highlight_color = Green
                    if r.text.strip():
                        Gtext.append(r.text)
                
                
                
                

                        
        paragraph = row.cells[2].paragraphs[0]
        for run in Ytext:
            nrun = paragraph.add_run(run + ';\n')
            nrun.font.highlight_color = Yellow
        for run in Gtext:
            nrun = paragraph.add_run(run + ';\n')
            nrun.font.highlight_color = Green
        for run in Ctext:
            nrun = paragraph.add_run(run + ';\n')
            nrun.font.highlight_color = Cyan
doc.save(name)